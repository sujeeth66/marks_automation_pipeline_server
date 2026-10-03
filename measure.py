#!/usr/bin/env python3
"""Measure how well a recognizer reads photographed marks sheets, using the typed workbook as the answer key.

Usage:
  python measure.py --photos page1.jpg page2.jpg --xlsx class5A_entry.xlsx --tab English [--recognizer mlp --model model.json] [--save-crops crops/]

--photos  one photo per printed page, in page order (page 1 first). Each must show the whole table.
Steps: find the table and straighten it (grid.py) -> cut out every marks cell -> recognize each cell
-> compare with the workbook tab. Prints exact-match accuracy, split by kind of value, plus sample errors.
To try another recognizer, add a function to RECOGNIZERS: it takes (grayscale cell image, max_value)
and returns (text, confidence); max_value is the largest mark the column allows (10 or 20).
Recognizers: "mlp" (the model trained by train_digits.py) and "rapidocr" (a free general OCR, for comparison).
"""
import argparse
import collections
import os
import sys

import cv2
import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))  # config.py lives one folder up
from config import SLOTS  # noqa: E402
from grid import cells  # noqa: E402


def cell_crops(photo, margin=6):
    img = cv2.imread(photo)
    if img is None:
        raise SystemExit(f"cannot read photo: {photo}")
    warped, hl, vl = cells(img)
    gaps = [b - a for a, b in zip(hl, hl[1:])]
    if gaps[0] > 1.4 * np.median(gaps):  # first gap is taller: that is the header row, drop it
        hl = hl[1:]
    if len(vl) < 5:
        raise SystemExit(f"{photo}: found only {len(vl)} vertical lines, expected 8 (is the whole table in the photo?)")
    cols = vl[-5:]  # the four mark columns
    gray = cv2.cvtColor(warped, cv2.COLOR_BGR2GRAY)
    return [[gray[r0 + margin:r1 - margin, c0 + margin:c1 - margin] for c0, c1 in zip(cols, cols[1:])]
            for r0, r1 in zip(hl, hl[1:])]


def rapidocr_recognizer(model=None):
    from rapidocr_onnxruntime import RapidOCR  # pip install rapidocr-onnxruntime
    engine = RapidOCR()

    def read(cell, max_value):
        c = cv2.copyMakeBorder(cell, 12, 12, 24, 24, cv2.BORDER_CONSTANT, value=255)
        c = cv2.resize(c, None, fx=64 / c.shape[0], fy=64 / c.shape[0], interpolation=cv2.INTER_CUBIC)
        res, _ = engine(cv2.cvtColor(c, cv2.COLOR_GRAY2BGR), use_det=False, use_cls=False, use_rec=True)
        return (res[0][0].strip().upper().replace(" ", ""), float(res[0][1])) if res else ("", 0.0)

    return read


def mlp_recognizer(model):
    from digits import decode, load_model, normalize, predict_proba, to_ink
    layers = load_model(model)

    def read(cell, max_value):
        vec, ink = normalize(to_ink(cell))
        if not ink:
            return "", 1.0  # nothing written
        return decode(predict_proba(layers, vec), max_value)

    return read


RECOGNIZERS = {"mlp": mlp_recognizer, "rapidocr": rapidocr_recognizer}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--photos", nargs="+", required=True)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--tab", required=True)
    ap.add_argument("--recognizer", default="mlp", choices=sorted(RECOGNIZERS))
    ap.add_argument("--model", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "model.json"),
                    help="model.json for the mlp recognizer")
    ap.add_argument("--save-crops")
    a = ap.parse_args()

    truth = []
    for row in load_workbook(a.xlsx, data_only=True)[a.tab].iter_rows(min_row=2, max_col=6, values_only=True):
        if row[0] is not None:
            truth.append([("" if v is None else str(v).strip().upper()) for v in row[2:6]])
    rows = [r for p in a.photos for r in cell_crops(p)]
    if len(rows) != len(truth):
        raise SystemExit(f"photos give {len(rows)} rows but the workbook tab has {len(truth)} students")

    read = RECOGNIZERS[a.recognizer](a.model)
    if a.save_crops:
        os.makedirs(a.save_crops, exist_ok=True)
    ok = total = 0
    kinds = collections.defaultdict(lambda: [0, 0])
    errors, all_reads = [], []
    for i, (crow, trow) in enumerate(zip(rows, truth), start=1):
        for j, (cell, t) in enumerate(zip(crow, trow), start=1):
            if a.save_crops:
                cv2.imwrite(os.path.join(a.save_crops, f"r{i:02d}_tool{j}.png"), cell)
            if t == "":
                continue
            text, conf = read(cell, SLOTS[j - 1][2])
            good = text == t
            kind = "A" if t == "A" else ("1 digit" if len(t) == 1 else "2 digits")
            kinds[kind][0] += good
            kinds[kind][1] += 1
            ok += good
            total += 1
            all_reads.append((i, j, t, text, conf))
            if not good:
                errors.append((i, j, t, text, round(conf, 2)))
    print(f"{total} cells, exact match {ok}/{total} = {100 * ok / total:.1f}%")
    print("by kind:", {k: f"{x}/{y}" for k, (x, y) in kinds.items()})
    print("sample errors (row, tool, typed, read, confidence):", errors[:12])
    wrong = collections.Counter((t, r) for _, _, t, r, _ in errors)
    print("most common mix-ups (typed -> read):", wrong.most_common(8))
    print("if cells above a confidence were accepted without a look:")
    for th in (0.9, 0.97, 0.99):
        above = [e for e in all_reads if e[4] >= th]
        print(f"  >= {th}: {len(above)} cells, {sum(1 for e in above if e[2] != e[3])} wrong")


if __name__ == "__main__":
    main()
