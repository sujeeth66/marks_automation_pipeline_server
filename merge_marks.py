#!/usr/bin/env python3
"""Merge the typed Excel workbook into one data file for fillMarks().

Usage:
  python merge_marks.py --class 3 --section A [--entry out/entry/class3A_entry.xlsx] [--working-days 37] [--out data_class3A.json]

Reads the workbook made by make_sheets.py: a tab per subject plus Attendance.
Matching is by student_id, never by row position.
Checks: marks are whole numbers within each slot's max, or A; attendance is a whole
number (and not above --working-days); every tab has the same students.
If anything is wrong it lists the problems and writes nothing. Blank cells are not
errors (they are listed as incomplete, and fillMarks() skips them).
"""
import argparse
import json
import re
import sys

from openpyxl import load_workbook

from config import CLASS_SUBJECTS, SLOTS, SUBJECT_LABELS


class WorkbookError(Exception):
    """The workbook is not in the shape make_sheets.py produced."""


def text(v):
    """Cell value -> clean string (Excel may hand back 10.0 or a number for an ID)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_rows(wb, sheet, ncols):
    """-> [(excel_row, student_id, name, [ncols cell strings])] for rows that have a student_id."""
    if sheet not in wb.sheetnames:
        raise WorkbookError(f'the workbook has no tab called "{sheet}"')
    out = []
    for r, row in enumerate(wb[sheet].iter_rows(min_row=2, max_col=2 + ncols, values_only=True), start=2):
        sid = text(row[0])
        if sid:
            out.append((r, sid, text(row[1]), [text(v).upper() for v in row[2:2 + ncols]]))
    return out


def parse_mark(v, mx):
    """-> (value, error). Blank gives ("", None)."""
    if v == "":
        return "", None
    if v == "A":
        return "A", None
    if re.fullmatch(r"\d{1,3}", v) and int(v) <= mx:
        return int(v), None
    return None, f'"{v}" is not a whole number from 0 to {mx}, or A'


def merge_workbook(wb, cls, working_days=None):
    """Check an opened entry workbook and merge it. -> (data, errors, incomplete).
    data is what fillMarks() reads. Nothing is raised for bad cells; they are listed in errors."""
    errors, data, names = [], {}, {}

    # attendance
    for r, sid, name, (v,) in read_rows(wb, "Attendance", 1):
        if sid in data:
            errors.append(f"Attendance row {r}: duplicate student_id {sid}")
            continue
        names[sid] = name
        v = re.sub(r"\.0+$", "", v)
        if v == "":
            att = ""
        elif re.fullmatch(r"\d{1,3}", v):
            att = int(v)
            if working_days and att > working_days:
                errors.append(f"Attendance row {r} ({name}): {att} is more than {working_days} working days")
        else:
            att = None
            errors.append(f'Attendance row {r} ({name}): "{v}" is not a whole number')
        data[sid] = {"attendance": att, "marks": {}}

    # subjects
    for key in CLASS_SUBJECTS[cls]:
        label = SUBJECT_LABELS[key]
        seen = set()
        for r, sid, name, cells in read_rows(wb, label, len(SLOTS)):
            if sid in seen:
                errors.append(f"{label} row {r}: duplicate student_id {sid}")
                continue
            seen.add(sid)
            if sid not in data:
                errors.append(f"{label} row {r}: student_id {sid} is not on the Attendance tab")
                continue
            vals = []
            for (slot_name, _, mx), v in zip(SLOTS, cells):
                val, err = parse_mark(v, mx)
                if err:
                    errors.append(f"{label} row {r} ({name}), {slot_name}: {err}")
                vals.append(val)
            data[sid]["marks"][key] = vals
        for sid in data:
            if sid not in seen:
                errors.append(f"{label}: student_id {sid} ({names[sid]}) is missing from this tab")

    incomplete = []
    for sid, rec in data.items():
        miss = []
        if rec["attendance"] == "":
            miss.append("attendance")
        for key, vals in rec["marks"].items():
            miss += [f"{SUBJECT_LABELS[key]} {SLOTS[i][0]}" for i, v in enumerate(vals) if v == ""]
        if miss:
            incomplete.append((sid, names[sid], miss))
    return data, errors, incomplete


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--class", dest="cls", type=int, required=True, choices=sorted(CLASS_SUBJECTS))
    ap.add_argument("--section", required=True)
    ap.add_argument("--entry")
    ap.add_argument("--working-days", type=int)
    ap.add_argument("--out")
    a = ap.parse_args()

    tag = f"class{a.cls}{a.section}"
    entry = a.entry or f"out/entry/{tag}_entry.xlsx"
    out_path = a.out or f"data_{tag}.json"
    try:
        wb = load_workbook(entry, data_only=True)
    except FileNotFoundError:
        sys.exit(f"missing file: {entry}")
    try:
        data, errors, incomplete = merge_workbook(wb, a.cls, a.working_days)
    except WorkbookError as e:
        sys.exit(str(e))

    if errors:
        print(f"{len(errors)} problem(s) found - nothing written:")
        for e in errors:
            print("  -", e)
        sys.exit(1)

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=1)

    print(f"{len(data)} students merged -> {out_path}")
    if incomplete:
        print(f"{len(incomplete)} student(s) have blank cells (the portal's Final Submit needs all of them filled):")
        for sid, name, miss in incomplete[:15]:
            print(f"  - {sid} {name}: {len(miss)} blank ({', '.join(miss[:4])}{'...' if len(miss) > 4 else ''})")
        if len(incomplete) > 15:
            print(f"  ... and {len(incomplete) - 15} more")
    print("Next: on the portal page, in the console, type  var data = <paste the JSON file's contents>  then  fillMarks(data)")


if __name__ == "__main__":
    main()
