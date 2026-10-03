"""Tests for the server. Run from the project folder:  python -m unittest discover -s tests -v"""
import io
import json
import os
import random
import sys
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ["API_KEYS"] = "test-key-aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa,second-key-bbbbbbbbbbbbbbbbbbbbbbbbbbbb"
os.environ["ALLOWED_ORIGINS"] = "chrome-extension://abcdef"

import app as server  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

KEY = {"X-API-Key": "test-key-aaaaaaaaaaaaaaaaaaaaaaaaaaaaaaaa"}
FIX = os.path.join(ROOT, "tests", "fixtures")


def roster(n=6):
    return [{"id": str(1000 + i), "name": f"STUDENT {i} & SON <X>"} for i in range(n)]


def body(**kw):
    b = {"class": 3, "section": "A", "school": "TEST SCHOOL & CO", "working_days": 37, "roster": roster()}
    b.update(kw)
    return b


class ServerTests(unittest.TestCase):
    def setUp(self):
        self.c = server.app.test_client()
        server._hits.clear()

    # ---- basics and auth ----
    def test_health_needs_no_key(self):
        self.assertEqual(self.c.get("/health").get_json(), {"ok": True})

    def test_wrong_or_missing_key_is_refused(self):
        self.assertEqual(self.c.post("/sheets", json=body()).status_code, 401)
        self.assertEqual(self.c.post("/sheets", json=body(), headers={"X-API-Key": "nope"}).status_code, 401)

    def test_second_key_works_and_responses_are_not_cached(self):
        r = self.c.post("/sheets", json=body(), headers={"X-API-Key": "second-key-bbbbbbbbbbbbbbbbbbbbbbbbbbbb"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["Cache-Control"], "no-store")

    def test_cors_only_for_listed_origins(self):
        ok = self.c.get("/health", headers={"Origin": "chrome-extension://abcdef"})
        bad = self.c.get("/health", headers={"Origin": "https://evil.example"})
        self.assertEqual(ok.headers.get("Access-Control-Allow-Origin"), "chrome-extension://abcdef")
        self.assertIsNone(bad.headers.get("Access-Control-Allow-Origin"))
        pre = self.c.options("/sheets", headers={"Origin": "chrome-extension://abcdef"})
        self.assertEqual(pre.status_code, 204)

    def test_rate_limit(self):
        old, server.RATE_PER_MIN = server.RATE_PER_MIN, 3
        try:
            codes = [self.c.post("/sheets", json={}, headers=KEY).status_code for _ in range(5)]
        finally:
            server.RATE_PER_MIN = old
        self.assertEqual(codes, [400, 400, 400, 429, 429])

    # ---- /sheets ----
    def test_sheets_zip_contents(self):
        r = self.c.post("/sheets", json=body(), headers=KEY)
        self.assertEqual(r.status_code, 200)
        names = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
        self.assertEqual(len(names), 4 + 1 + 1)  # class 3: 4 subject PDFs + attendance PDF + workbook
        self.assertIn("entry/class3A_entry.xlsx", names)
        self.assertIn("sheets/class3A_Maths.pdf", names)

    def test_classes_1_and_2_have_three_subjects(self):
        r = self.c.post("/sheets", json=body(**{"class": 1}), headers=KEY)
        names = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
        self.assertEqual(len(names), 3 + 1 + 1)
        self.assertNotIn("sheets/class1A_EVS.pdf", names)

    def test_sheets_rejects_bad_input(self):
        for bad in (body(**{"class": 9}), body(section=""), body(section="TOOLONG"), body(working_days=0),
                    body(roster=[]), body(roster=roster(121)), body(roster=[{"id": "../x", "name": "A"}]),
                    body(roster=[{"id": "1", "name": ""}]), body(roster=[{"id": "1", "name": "A"}] * 2)):
            self.assertEqual(self.c.post("/sheets", json=bad, headers=KEY).status_code, 400, bad.get("section"))
        self.assertEqual(self.c.post("/sheets", data="not json", headers=KEY).status_code, 400)

    # ---- /merge ----
    def _workbook(self, mutate=None):
        z = zipfile.ZipFile(io.BytesIO(self.c.post("/sheets", json=body(), headers=KEY).data))
        wb = load_workbook(io.BytesIO(z.read("entry/class3A_entry.xlsx")))
        random.seed(1)
        for ws in wb.worksheets[1:]:
            for r in range(2, ws.max_row + 1):
                if ws.title == "Attendance":
                    ws.cell(row=r, column=3, value=random.randint(20, 37))
                else:
                    for c, mx in zip(range(3, 7), (10, 10, 10, 20)):
                        ws.cell(row=r, column=c, value="A" if random.random() < 0.1 else random.randint(0, mx))
        if mutate:
            mutate(wb)
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return out

    def _merge(self, wb, **form):
        data = {"workbook": (wb, "w.xlsx"), "class": "3", "working_days": "37"}
        data.update(form)
        return self.c.post("/merge", data=data, headers=KEY, content_type="multipart/form-data")

    def test_round_trip_gives_data_for_fillMarks(self):
        r = self._merge(self._workbook())
        self.assertEqual(r.status_code, 200)
        j = r.get_json()
        self.assertTrue(j["ok"])
        self.assertEqual(sorted(j["data"]), [str(1000 + i) for i in range(6)])
        rec = j["data"]["1000"]
        self.assertEqual(set(rec["marks"]), {"1stLang", "Eng", "Math", "ES"})
        self.assertEqual(len(rec["marks"]["Eng"]), 4)
        self.assertIsInstance(rec["attendance"], int)

    def test_bad_cells_are_reported_and_nothing_is_returned(self):
        def mutate(wb):
            wb["Maths"]["C2"] = "x"
            wb["Maths"]["F3"] = 21
            wb["Attendance"]["C4"] = 99
        r = self._merge(self._workbook(mutate))
        self.assertEqual(r.status_code, 422)
        j = r.get_json()
        self.assertFalse(j["ok"])
        self.assertEqual(len(j["errors"]), 3)
        self.assertNotIn("data", j)

    def test_merge_rejects_wrong_uploads(self):
        self.assertEqual(self._merge(io.BytesIO(b"not a zip")).status_code, 400)
        self.assertEqual(self.c.post("/merge", data={"class": "3"}, headers=KEY,
                                     content_type="multipart/form-data").status_code, 400)
        def drop_tab(wb):
            del wb["EVS"]
        self.assertEqual(self._merge(self._workbook(drop_tab)).status_code, 400)

    # ---- /scan ----
    def _scan(self, files, **form):
        data = {"class": "5", "tab": "Eng"}
        data.update(form)
        data["photos"] = [(open(os.path.join(FIX, f), "rb"), f) for f in files]
        return self.c.post("/scan", data=data, headers=KEY, content_type="multipart/form-data")

    def test_scan_reads_the_sample_sheet(self):
        r = self._scan(["english_page1.jpg", "english_page2.jpg"], expected_rows="38")
        self.assertEqual(r.status_code, 200)
        rows = r.get_json()["rows"]
        self.assertEqual(len(rows), 38)
        self.assertTrue(all(len(row) == 4 for row in rows))
        truth = json.load(open(os.path.join(FIX, "english_truth.json")))
        ok = sum(c["value"] == truth[i][j] for i, row in enumerate(rows) for j, c in enumerate(row))
        self.assertGreater(ok / 152, 0.8, f"only {ok}/152 cells matched")
        self.assertTrue(rows[0][0]["png"].startswith("iVBOR"))  # a PNG

    def test_scan_checks_row_count_and_inputs(self):
        self.assertEqual(self._scan(["english_page1.jpg"], expected_rows="38").status_code, 422)
        self.assertEqual(self._scan(["english_page1.jpg"], tab="Nonsense").status_code, 400)
        self.assertEqual(self._scan(["english_page1.jpg"], **{"class": "7"}).status_code, 400)
        self.assertEqual(self.c.post("/scan", data={"class": "5", "tab": "Eng"}, headers=KEY,
                                     content_type="multipart/form-data").status_code, 400)

    def test_scan_rejects_non_images(self):
        data = {"class": "5", "tab": "Eng", "photos": (io.BytesIO(b"hello"), "x.jpg")}
        r = self.c.post("/scan", data=data, headers=KEY, content_type="multipart/form-data")
        self.assertEqual(r.status_code, 400)
        self.assertIn("photo 1", r.get_json()["error"])

    def test_oversize_upload_is_refused(self):
        old = server.app.config["MAX_CONTENT_LENGTH"]
        server.app.config["MAX_CONTENT_LENGTH"] = 1000
        try:
            r = self.c.post("/scan", data={"class": "5", "tab": "Eng", "photos": (io.BytesIO(b"x" * 5000), "x.jpg")},
                            headers=KEY, content_type="multipart/form-data")
        finally:
            server.app.config["MAX_CONTENT_LENGTH"] = old
        self.assertEqual(r.status_code, 413)

    def test_attendance_scan_does_not_allow_A(self):
        # the sample sheet has 4 mark columns, so the attendance layout (1 column) reads the last one
        r = self._scan(["english_page1.jpg"], tab="Attendance", working_days="37")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(all(c["value"] != "A" for row in r.get_json()["rows"] for c in row))


if __name__ == "__main__":
    unittest.main()
